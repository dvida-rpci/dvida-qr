"""Lógica pura del plan de mantenimiento: biblioteca sabana_base.xlsx, sugerencia de
clase por TAG y detección de pares A/B.

No depende de Supabase ni de la GUI (la usa gui_plan.py en la fase 2).
Diseño: docs/superpowers/specs/2026-10-14-plan-mantenimiento-design.md
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass

import openpyxl
from openpyxl.styles import Font, PatternFill

TAG_TYPE_RE = re.compile(r"^\d{3}-([A-Z]+)-")
PAIR_RE = re.compile(r"^(.*\d)([A-Z])$")
SUGGEST_FIELDS = ("TIPO DE BOMBA", "TIPO", "SERVICIO", "Título")
CLASS_HEADERS = ["code", "nombre", "horas_dia", "letras_tag", "palabras_clave"]
ACTIVITY_HEADERS = ["clase_code", "code", "nombre", "cada_horas", "cada_semanas"]


class LibraryError(ValueError):
    """sabana_base.xlsx inválida; `errors` trae un mensaje por problema (con hoja y fila)."""

    def __init__(self, errors: list[str]):
        super().__init__("\n".join(errors))
        self.errors = errors


@dataclass
class EquipmentClass:
    code: str
    name: str
    hours_per_day: float
    tag_letters: list[str]
    keywords: list[str]


@dataclass
class Activity:
    class_code: str
    code: str
    name: str
    interval_hours: float | None
    interval_weeks: int | None


@dataclass
class ImportPlan:
    new_classes: list[EquipmentClass]
    new_activities: list[Activity]


def normalize(text) -> str:
    """Minúsculas y sin tildes, para comparar palabras clave."""
    decomposed = unicodedata.normalize("NFKD", str(text or ""))
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch)).lower().strip()


def _text(value) -> str:
    return str(value).strip() if value is not None else ""


def _split(value) -> list[str]:
    return [part.strip() for part in _text(value).split(";") if part.strip()]


def _rows(ws):
    """(número de fila Excel, dict encabezado→valor) por cada fila no vacía."""
    rows = ws.iter_rows(values_only=True)
    headers = [_text(h) for h in next(rows, [])]
    for row_num, values in enumerate(rows, start=2):
        if all(_text(v) == "" for v in values):
            continue
        yield row_num, dict(zip(headers, values))


def _number(value, label: str, where: str, errors: list[str], *, integer: bool = False):
    """Número > 0, o None si la celda está vacía. Los problemas se acumulan en `errors`."""
    if _text(value) == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        errors.append(f"{where}: {label} debe ser un número (es «{value}»)")
        return None
    if number <= 0:
        errors.append(f"{where}: {label} debe ser mayor que 0")
        return None
    if integer:
        if number != int(number):
            errors.append(f"{where}: {label} debe ser un número entero")
            return None
        return int(number)
    return number


def load_library(path) -> tuple[list[EquipmentClass], list[Activity]]:
    """Lee y valida la biblioteca. Si hay cualquier error, no devuelve nada a medias:
    lanza LibraryError con todos los problemas encontrados."""
    wb = openpyxl.load_workbook(path, data_only=True)
    missing = [name for name in ("Clases", "Actividades") if name not in wb.sheetnames]
    if missing:
        raise LibraryError([f"Falta la hoja «{name}»" for name in missing])

    errors: list[str] = []
    classes: list[EquipmentClass] = []
    for row_num, row in _rows(wb["Clases"]):
        where = f"Clases fila {row_num}"
        before = len(errors)
        code = _text(row.get("code"))
        name = _text(row.get("nombre"))
        letters = _split(row.get("letras_tag"))
        hours = None
        if not code:
            errors.append(f"{where}: falta code")
        elif any(c.code == code for c in classes):
            errors.append(f"{where}: code «{code}» repetido")
        if not name:
            errors.append(f"{where}: falta nombre")
        if _text(row.get("horas_dia")) == "":
            errors.append(f"{where}: falta horas_dia")
        else:
            hours = _number(row.get("horas_dia"), "horas_dia", where, errors)
            if hours is not None and hours > 24:
                errors.append(f"{where}: horas_dia no puede pasar de 24")
        if not letters:
            errors.append(f"{where}: falta letras_tag")
        if len(errors) == before:
            classes.append(
                EquipmentClass(code, name, hours, letters, _split(row.get("palabras_clave")))
            )

    class_codes = {c.code for c in classes}
    activities: list[Activity] = []
    for row_num, row in _rows(wb["Actividades"]):
        where = f"Actividades fila {row_num}"
        before = len(errors)
        class_code = _text(row.get("clase_code"))
        code = _text(row.get("code"))
        name = _text(row.get("nombre"))
        hours = _number(row.get("cada_horas"), "cada_horas", where, errors)
        weeks = _number(row.get("cada_semanas"), "cada_semanas", where, errors, integer=True)
        if class_code not in class_codes:
            errors.append(f"{where}: clase_code «{class_code}» no existe en la hoja Clases")
        if not code:
            errors.append(f"{where}: falta code")
        elif any(a.class_code == class_code and a.code == code for a in activities):
            errors.append(f"{where}: code «{code}» repetido en la clase {class_code}")
        if not name:
            errors.append(f"{where}: falta nombre")
        if hours is None and weeks is None and len(errors) == before:
            errors.append(f"{where}: indica cada_horas, cada_semanas o ambas")
        if len(errors) == before:
            activities.append(Activity(class_code, code, name, hours, weeks))

    if errors:
        raise LibraryError(errors)
    return classes, activities


def plan_import(
    classes: list[EquipmentClass],
    activities: list[Activity],
    existing_class_codes: set[str],
    existing_activity_keys: set[tuple[str, str]],
) -> ImportPlan:
    """Merge por code: solo lo que la planta todavía no tiene. Lo existente no se toca,
    así se preservan las frecuencias ya ajustadas en la planta."""
    return ImportPlan(
        new_classes=[c for c in classes if c.code not in existing_class_codes],
        new_activities=[
            a for a in activities if (a.class_code, a.code) not in existing_activity_keys
        ],
    )


def suggest_class(tag: str, props: dict, classes: list[EquipmentClass]) -> str | None:
    """Primera clase (en el orden de la hoja) cuya letra de TAG coincide y alguna de cuyas
    palabras clave aparece en TIPO DE BOMBA / TIPO / SERVICIO / Título."""
    match = TAG_TYPE_RE.match(tag or "")
    if not match:
        return None
    letter = match.group(1)
    text = " ".join(normalize(props.get(field)) for field in SUGGEST_FIELDS)
    for cls in classes:
        if letter in cls.tag_letters and any(normalize(k) in text for k in cls.keywords):
            return cls.code
    return None


def detect_pairs(tags: list[str]) -> dict[str, str]:
    """{tag: group_key} para TAGs que solo difieren en la letra final (100-P-01A/B).
    Un TAG con letra final pero sin hermano no forma par."""
    groups: dict[str, list[str]] = {}
    for tag in tags:
        match = PAIR_RE.match(tag)
        if match:
            groups.setdefault(match.group(1), []).append(tag)
    return {tag: prefix for prefix, members in groups.items() if len(members) >= 2 for tag in members}


def load_plantilla_tags(path) -> list[dict]:
    """Filas de la hoja Datos de plantilla_sitio.xlsx que tienen TAG, como dicts."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return [row for _, row in _rows(wb["Datos"]) if _text(row.get("TAG"))]


def _write_sheet(ws, headers: list[str], rows: list[list]):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C3AED")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"


def export_library(classes: list[EquipmentClass], activities: list[Activity], path) -> None:
    """Escribe la sábana en el formato de la biblioteca (lo inverso de load_library)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clases"
    _write_sheet(ws, CLASS_HEADERS, [
        [c.code, c.name, c.hours_per_day, ";".join(c.tag_letters), ";".join(c.keywords)]
        for c in classes
    ])
    _write_sheet(wb.create_sheet("Actividades"), ACTIVITY_HEADERS, [
        [a.class_code, a.code, a.name, a.interval_hours, a.interval_weeks] for a in activities
    ])
    wb.save(path)
