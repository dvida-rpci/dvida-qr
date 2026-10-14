"""maintenance_plan.py: biblioteca sabana_base.xlsx, sugerencia de clase y pares (sin red)."""
import openpyxl
import pytest

from maintenance_plan import (
    Activity,
    EquipmentClass,
    LibraryError,
    detect_pairs,
    export_library,
    load_library,
    load_plantilla_tags,
    normalize,
    plan_import,
    suggest_class,
)

CLASS_HEADERS = ["code", "nombre", "horas_dia", "letras_tag", "palabras_clave"]
ACTIVITY_HEADERS = ["clase_code", "code", "nombre", "cada_horas", "cada_semanas"]


def _write_library(path, clases, actividades=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clases"
    ws.append(CLASS_HEADERS)
    for row in clases:
        ws.append(row)
    if actividades is not None:
        ws = wb.create_sheet("Actividades")
        ws.append(ACTIVITY_HEADERS)
        for row in actividades:
            ws.append(row)
    wb.save(path)
    return path


def _cls(code, letters=("P",), keywords=()):
    return EquipmentClass(code, f"Clase {code}", 24, list(letters), list(keywords))


def test_normalize_quita_tildes_y_mayusculas():
    assert normalize(" Centrífuga ") == "centrifuga"
    assert normalize(None) == ""


def test_biblioteca_del_repo_trae_el_ejemplo_de_centrifugas():
    classes, activities = load_library("sabana_base.xlsx")
    assert [c.code for c in classes] == ["BOMBA_CENTRIFUGA_H"]
    assert classes[0].tag_letters == ["P"]
    assert {(a.code, a.interval_hours, a.interval_weeks) for a in activities} == {
        ("LUB-ROD", 1300, 12),
        ("REV-SELLO", 4000, 26),
        ("ALINEACION", 8000, 52),
    }


def test_biblioteca_invalida_reporta_cada_error_con_su_fila(tmp_path):
    path = _write_library(
        tmp_path / "s.xlsx",
        [
            ["A", "Clase A", 24, "P", "centrifuga"],
            ["A", "Repetida", 8, "P", ""],
            ["B", "", 30, "", ""],
        ],
        [
            ["NOEXISTE", "X", "Algo", 100, None],
            ["A", "Y", "Sin frecuencia", None, None],
            ["A", "Z", "Semanas raras", None, 1.5],
            ["A", "W", "Horas negativas", -5, None],
        ],
    )
    with pytest.raises(LibraryError) as exc:
        load_library(path)
    errors = exc.value.errors
    assert "Clases fila 3: code «A» repetido" in errors
    assert "Clases fila 4: falta nombre" in errors
    assert "Clases fila 4: horas_dia no puede pasar de 24" in errors
    assert "Clases fila 4: falta letras_tag" in errors
    assert "Actividades fila 2: clase_code «NOEXISTE» no existe en la hoja Clases" in errors
    assert "Actividades fila 3: indica cada_horas, cada_semanas o ambas" in errors
    assert "Actividades fila 4: cada_semanas debe ser un número entero" in errors
    assert "Actividades fila 5: cada_horas debe ser mayor que 0" in errors


def test_biblioteca_sin_hoja_actividades(tmp_path):
    path = _write_library(tmp_path / "s.xlsx", [["A", "Clase A", 24, "P", ""]])
    with pytest.raises(LibraryError) as exc:
        load_library(path)
    assert exc.value.errors == ["Falta la hoja «Actividades»"]


def test_plan_import_solo_agrega_lo_que_falta():
    classes = [_cls("A"), _cls("B")]
    activities = [
        Activity("A", "L", "Lubricación", 1300, 12),
        Activity("A", "S", "Sello", 4000, None),
        Activity("B", "X", "Otra", None, 4),
    ]
    plan = plan_import(classes, activities, {"A"}, {("A", "L")})
    assert [c.code for c in plan.new_classes] == ["B"]
    assert [(a.class_code, a.code) for a in plan.new_activities] == [("A", "S"), ("B", "X")]


ROWS = [
    {"TAG": "200-P-01A", "TIPO DE BOMBA": "Centrifuga horizontal", "SERVICIO": "Bomba centrifuga horizontal"},
    {"TAG": "500-P-01", "TIPO DE BOMBA": "Centrifuga horizontal", "SERVICIO": "Bomba de trasiego de agua filtrada"},
    {"TAG": "802-P-01", "TIPO": "Neumática", "SERVICIO": "Bomba dosificadora de floculante"},
    {"TAG": "801-P-01", "TIPO": "Diafragma eléctrica", "SERVICIO": "Bomba dosificadora"},
    {"TAG": "100-TK-01", "SERVICIO": "Homogenización centrífuga"},
    {"TAG": "Hoja3", "SERVICIO": "centrifuga"},
]


def test_suggest_class_con_la_biblioteca_del_repo():
    classes, _ = load_library("sabana_base.xlsx")
    sugeridas = {row["TAG"]: suggest_class(row["TAG"], row, classes) for row in ROWS}
    assert sugeridas == {
        "200-P-01A": "BOMBA_CENTRIFUGA_H",
        "500-P-01": "BOMBA_CENTRIFUGA_H",
        "802-P-01": None,
        "801-P-01": None,
        "100-TK-01": None,
        "Hoja3": None,
    }


def test_suggest_class_gana_la_primera_clase_de_la_hoja():
    classes = [_cls("PRIMERA", keywords=["bomba"]), _cls("SEGUNDA", keywords=["dosificadora"])]
    assert suggest_class("801-P-01", {"SERVICIO": "Bomba dosificadora"}, classes) == "PRIMERA"


def test_detect_pairs():
    tags = [
        "100-P-01A", "100-P-01B", "500-P-01", "400-MX-01", "400-MX-02",
        "400-C-01A", "400-C-01B", "802-P-01A",
    ]
    assert detect_pairs(tags) == {
        "100-P-01A": "100-P-01",
        "100-P-01B": "100-P-01",
        "400-C-01A": "400-C-01",
        "400-C-01B": "400-C-01",
    }


def test_load_plantilla_tags(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["TAG", "Categoría", "SERVICIO"])
    ws.append(["100-P-01A", "EQUIPOS", "Bomba centrifuga horizontal"])
    ws.append([None, None, None])
    ws.append([None, "EQUIPOS", "fila sin TAG"])
    ws.append(["100-TK-01", "TANQUES", "Homogenización"])
    wb.create_sheet("Instrucciones")
    wb.save(tmp_path / "plantilla.xlsx")
    rows = load_plantilla_tags(tmp_path / "plantilla.xlsx")
    assert [r["TAG"] for r in rows] == ["100-P-01A", "100-TK-01"]
    assert rows[0]["SERVICIO"] == "Bomba centrifuga horizontal"


def test_export_library_ida_y_vuelta(tmp_path):
    classes, activities = load_library("sabana_base.xlsx")
    export_library(classes, activities, tmp_path / "exportada.xlsx")
    assert load_library(tmp_path / "exportada.xlsx") == (classes, activities)
